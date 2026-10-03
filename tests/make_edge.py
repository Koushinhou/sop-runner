# 边界用例：无样式手工编号、粗体手工标题、项目符号、合并单元格、英文 expected
from docx import Document
from docx.shared import Pt
d = Document()
def bold(t):
    p = d.add_paragraph(); p.add_run(t).bold = True
bold('1. 目的')
d.add_paragraph('DBサーバ db01.example.local（192.0.2.21）の設定変更を行う。')
bold('2. 手順')
d.add_paragraph('(1) 設定ファイルをバックアップする。')
d.add_paragraph('[root@db01 ~]# cp -p /etc/my.cnf /etc/my.cnf.bak')
d.add_paragraph('Expected: バックアップファイルが作成されること')
d.add_paragraph('(2) 設定値を確認する。')
d.add_paragraph('・max_connections の値：（　　）')
d.add_paragraph('・バッファサイズ：【　】MB')
d.add_paragraph('(3) サービスを再起動する。')
d.add_paragraph('systemctl restart mysqld')
d.add_paragraph('確認結果：')
bold('3. 確認')
t = d.add_table(rows=4, cols=4); t.style = 'Table Grid'
for j, h in enumerate(['区分', '項目', '確認内容', 'チェック']): t.cell(0, j).text = h
t.cell(1, 0).merge(t.cell(2, 0)); t.cell(1, 0).text = 'DB'
t.cell(1, 1).text = '接続確認'; t.cell(1, 2).text = '接続できること'
t.cell(2, 1).text = 'ログ確認'; t.cell(2, 2).text = 'エラーなし'
a = t.cell(3, 0).merge(t.cell(3, 2)); a.text = '総合判定'
d.save('/workspace/sop-runner/tests/edge.docx')
