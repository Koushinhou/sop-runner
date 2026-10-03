# 用 python-docx 校验导出的执行记录
import json, re, sys
from docx import Document
d = Document('tests/out/record.docx')
st = json.load(open('tests/out/state.json'))
ok = fail = 0
def check(c, m):
    global ok, fail
    print(('  ✔ ' if c else '  ✘ ') + m); ok += c; fail += (not c)
paras = [p.text for p in d.paragraphs]
check(paras[0] == '作業実施記録', 'タイトル: ' + paras[0])
check(len(d.tables) == 2, '表 2 個（ヘッダ情報 + 手順表）')
info = {r.cells[0].text: r.cells[1].text for r in d.tables[0].rows}
check(info.get('実施者') == '山田 太郎（テスト）', '実施者 = ' + info.get('実施者', ''))
check('Webサーバ定期パッチ適用手順書' in info.get('手順書名', ''), '手順書名 = ' + info.get('手順書名', ''))
ts = re.compile(r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$')
check(bool(ts.match(info.get('開始日時', ''))) and bool(ts.match(info.get('終了日時', ''))), '開始/終了日時 ' + info.get('開始日時', '') + ' → ' + info.get('終了日時', ''))
check(info.get('タイムゾーン') == 'UTC+09:00', 'タイムゾーン ' + info.get('タイムゾーン', ''))
t = d.tables[1]
hdr = [c.text for c in t.rows[0].cells]
check(hdr == ['No.', '手順', '期待結果', '記録値', '確認日時', '備考', '異常'], '列: ' + ' | '.join(hdr))
rows = t.rows[1:]
check(len(rows) == 18, '手順行 = %d' % len(rows))
check(all(ts.match(r.cells[4].text) for r in rows), '全行に確認日時（ローカル時刻）')
r9 = rows[8].cells
check(r9[6].text == 'あり' and 'タイムアウト' in r9[5].text and 'NG' in r9[3].text, '手順9: 異常=あり, 備考, 記録値 NG …')
check(sum(r.cells[6].text == 'あり' for r in rows) == 1, '異常は 1 件のみ')
check('適用対象パッケージ数：12 件' in rows[7].cells[3].text, '手順8 記録値: ' + rows[7].cells[3].text)
check('カーネルバージョン：5.14.0-999.el9.x86_64' in rows[9].cells[3].text, '手順10 記録値')
check('作業者：山田 太郎' in rows[0].cells[3].text and '確認者：佐藤 花子' in rows[0].cells[3].text, '手順1 記録値')
check('☑' in rows[17].cells[3].text, '手順18 チェック ☑')
check('lbctl disable --pool web --member 192.0.2.11' in rows[6].cells[1].text and 'セッション数 0' in rows[6].cells[5].text, '手順7 コマンド + 備考')
check('Complete!' in rows[8].cells[2].text, '期待結果列')
mono = [r for p in rows[6].cells[1].paragraphs for r in p.runs if r.font.name == 'Consolas']
check(len(mono) >= 1, 'コマンドは Consolas 等幅フォント')
check(d.sections[0].orientation == 1, '横向き A4')
print('\nRESULT: %d passed, %d failed' % (ok, fail)); sys.exit(1 if fail else 0)
