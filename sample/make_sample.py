# -*- coding: utf-8 -*-
"""架空の手順書サンプルを生成する（実在の企業・サーバ情報は一切含まない）。"""
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import os

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Yu Gothic'
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')

def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.5)
    p.paragraph_format.left_indent = Pt(24)
    return p

def para(text, indent=True):
    p = doc.add_paragraph(text)
    if indent:
        p.paragraph_format.left_indent = Pt(24)
    return p

def step(text):
    return doc.add_paragraph(text, style='List Number')

doc.add_paragraph('Webサーバ定期パッチ適用手順書', style='Title')
doc.add_paragraph('文書番号：OPS-WEB-0001（架空）　版数：1.0　※本書はサンプル用の架空の手順書です。')

t = doc.add_table(rows=3, cols=2)
t.style = 'Table Grid'
for i, k in enumerate(['作業日', '作業者', '確認者']):
    t.cell(i, 0).text = k

doc.add_heading('1. 概要', level=1)
doc.add_paragraph('本手順書は、web01.example.local および web02.example.local（192.0.2.11 / 192.0.2.12）に対して'
                  '月次セキュリティパッチを適用する手順を示す。作業は計画停止時間内に実施すること。')

doc.add_heading('2. 事前準備', level=1)
step('作業端末から踏み台サーバ（192.0.2.10）へSSH接続する。')
code('$ ssh opeuser@192.0.2.10')
para('期待結果：踏み台サーバのプロンプト [opeuser@bastion01 ~]$ が表示されること。')
step('作業開始時刻を記入する。')
para('開始時刻：＿＿＿＿＿＿')
step('対象サーバのバックアップが取得済みであることを確認する。')
para('確認結果：（　　）　※取得済みなら「OK」と記入')

doc.add_heading('3. パッチ適用作業', level=1)
doc.add_heading('3.1 web01.example.local', level=2)
step('web01へSSH接続し、root権限に切り替える。')
code('$ ssh opeuser@192.0.2.11')
code('$ sudo -i')
para('期待結果：プロンプトが [root@web01 ~]# になること。')
step('ロードバランサからweb01を切り離す。')
code('# lbctl disable --pool web --member 192.0.2.11')
para('期待結果：「member disabled」と表示されること。')
para('切り離し時刻：___________')
p = step('適用対象パッチを確認する（コマンド ')
r = p.add_run('dnf check-update --security'); r.font.name = 'Courier New'
p.add_run(' を実行）。')
para('適用対象パッケージ数：（　　）件')
step('パッチを適用する。')
code('# dnf update -y --security')
para('期待結果：最後に「Complete!」と表示されること。')
para('結果：')
step('サーバを再起動し、再接続後にカーネルバージョンを記入する。')
code('# shutdown -r now')
code('# uname -r')
para('カーネルバージョン（記入）：')

doc.add_heading('3.2 web02.example.local', level=2)
step('3.1 の手順を web02（192.0.2.12）に対して同様に実施する。')
code('$ ssh opeuser@192.0.2.12')
para('実施結果：＿＿＿＿')

doc.add_heading('4. 動作確認', level=1)
doc.add_paragraph('以下の確認項目を実施し、確認結果欄に OK / NG を記入すること。')
rows = [
    ('No', '確認項目', 'コマンド', '期待結果', '確認結果'),
    ('1', 'HTTP応答確認（web01）', 'curl -I http://192.0.2.11/', 'HTTP/1.1 200 OK', ''),
    ('2', 'HTTP応答確認（web02）', 'curl -I http://192.0.2.12/', 'HTTP/1.1 200 OK', ''),
    ('3', 'httpdプロセス確認', 'systemctl is-active httpd', 'active と表示されること', ''),
    ('4', 'エラーログ確認', 'tail -n 50 /var/log/httpd/error_log', '新たなエラーが出力されていないこと', ''),
]
t = doc.add_table(rows=len(rows), cols=5)
t.style = 'Table Grid'
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = t.cell(i, j)
        c.text = v
        if j == 2 and i > 0:
            for run in c.paragraphs[0].runs:
                run.font.name = 'Courier New'

doc.add_heading('5. 作業完了', level=1)
step('ロードバランサへweb01/web02を再組み込みする。')
code('# lbctl enable --pool web --member 192.0.2.11')
code('# lbctl enable --pool web --member 192.0.2.12')
para('期待結果：両メンバーが「active」となること。')
step('作業終了時刻を記入し、作業管理表へ記録する。')
para('終了時刻：＿＿＿＿＿＿')
step('作業結果を作業責任者へ報告する。')
para('□ 作業責任者へ報告した')

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Webサーバ定期パッチ適用手順書.docx')
doc.save(out)
print(out)
